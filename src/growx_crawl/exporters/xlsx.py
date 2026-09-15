from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from growx_crawl.enrichment.bde_brief import BDEBriefGenerator
from growx_crawl.models import Company
from growx_crawl.storage import ErrorRepository, JobRepository, LeadRepository


class XLSXExporter:
    def __init__(self, db_path: str = None):
        self.db_path = db_path
        self.job_repo = JobRepository(db_path)
        self.lead_repo = LeadRepository(db_path)
        self.error_repo = ErrorRepository(db_path)

    def export_job(self, job_id: str, output_path: str, stage: str = "all") -> str:
        job = self.job_repo.get_job(job_id)
        if not job:
            raise ValueError(f"Job not found: {job_id}")

        leads = self.lead_repo.get_job_lead_candidates(job_id)
        dedupes = self.lead_repo.get_job_dedupe_records(job_id)
        errors = self.error_repo.list_job_errors(job_id)
        evidence = self.lead_repo.get_job_evidence(job_id)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        # 1. Leads Sheet
        ws_leads = wb.create_sheet(title="Leads" if stage == "all" else "Discovered Candidates")
        self._write_leads_sheet(ws_leads, leads)

        # 2. Contacts Sheet
        ws_contacts = wb.create_sheet(title="Contacts")
        self._write_contacts_sheet(ws_contacts, leads)

        # 3. Duplicates Sheet
        ws_duplicates = wb.create_sheet(title="Duplicates")
        self._write_duplicates_sheet(ws_duplicates, dedupes)

        # 4. Evidence Sheet
        ws_evidence = wb.create_sheet(title="Evidence")
        self._write_evidence_sheet(ws_evidence, evidence)

        # 5. Errors Sheet
        ws_errors = wb.create_sheet(title="Errors")
        self._write_errors_sheet(ws_errors, errors)

        # 6. Run Summary Sheet
        ws_summary = wb.create_sheet(title="Run Summary")
        self._write_summary_sheet(ws_summary, job, len(leads), len(dedupes), len(errors))

        # Format worksheets
        for ws in wb.worksheets:
            self._apply_sheet_formatting(ws)

        out_file = Path(output_path)
        out_file.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out_file)
        return str(out_file)

    def _write_leads_sheet(self, ws, leads: List[Dict[str, Any]]):
        headers = [
            "Company", "Industry", "Category", "City", "State", "Country",
            "Website", "Website Missing", "Primary Phone", "Primary Email", "WhatsApp",
            "Instagram", "LinkedIn", "Primary Decision Maker", "Decision Maker Role",
            "Decision Maker Tier", "Discovered From", "Lead Score", "Priority",
            "Research Summary", "Collected At"
        ]
        ws.append(headers)

        for item in leads:
            company = Company(**item)
            brief = BDEBriefGenerator.generate_brief(company)
            disc_from = ", ".join(item.get("discovered_from", [])) or item.get("source", "")

            ws.append([
                brief.company_name,
                brief.industry or "",
                brief.category or "",
                brief.city or "",
                brief.state or "",
                brief.country or "",
                brief.website or "",
                "YES" if brief.website_missing else "NO",
                brief.primary_phone or "",
                brief.primary_email or "",
                brief.whatsapp or "",
                brief.instagram or "",
                brief.linkedin or "",
                brief.primary_decision_maker or "N/A",
                brief.decision_maker_role or "N/A",
                brief.decision_maker_tier or "N/A",
                disc_from,
                item.get("score", 0),
                item.get("priority", "Low"),
                brief.research_summary,
                item.get("created_at", ""),
            ])

    def _write_contacts_sheet(self, ws, leads: List[Dict[str, Any]]):
        headers = ["Company", "Contact Name", "Role/Title", "Decision Maker Tier", "Email", "Phone", "LinkedIn", "Source URL"]
        ws.append(headers)

        for item in leads:
            company_name = item.get("name", "")
            for cnt in item.get("contacts", []):
                ws.append([
                    company_name,
                    cnt.get("name", ""),
                    cnt.get("title", ""),
                    cnt.get("decision_maker_tier", "Tier 3"),
                    cnt.get("email", ""),
                    cnt.get("phone", ""),
                    cnt.get("linkedin_url", ""),
                    cnt.get("source_url", ""),
                ])

    def _write_duplicates_sheet(self, ws, dedupes: List[Dict[str, Any]]):
        headers = ["Canonical Company", "Duplicate Company", "Match Signals", "Confidence", "Decision", "Detected At"]
        ws.append(headers)

        for d in dedupes:
            ws.append([
                d.get("canonical_name", d.get("canonical_company_id", "")),
                d.get("duplicate_name", d.get("duplicate_company_id", "")),
                d.get("match_signals", ""),
                d.get("confidence", 0.0),
                d.get("decision", ""),
                d.get("created_at", ""),
            ])

    def _write_evidence_sheet(self, ws, evidence: List[Dict[str, Any]]):
        headers = ["Company", "Fact Key", "Extracted Value", "Extraction Type", "Source URL", "Confidence", "Timestamp"]
        ws.append(headers)

        for ev in evidence:
            ws.append([
                ev.get("company_name", ""),
                ev.get("key", ""),
                str(ev.get("value", "")),
                ev.get("extraction_type", "text"),
                ev.get("source_url", ""),
                ev.get("confidence", 1.0),
                ev.get("observed_at", ""),
            ])

    def _write_errors_sheet(self, ws, errors: List[Dict[str, Any]]):
        headers = ["Target URL", "Error Type", "Error Message", "Timestamp"]
        ws.append(headers)

        for err in errors:
            ws.append([
                err.get("url", ""),
                err.get("error_type", ""),
                err.get("message", ""),
                err.get("created_at", ""),
            ])

    def _write_summary_sheet(self, ws, job, lead_count: int, dup_count: int, error_count: int):
        ws.append(["Field", "Value"])
        ws.append(["Job ID", job.id])
        ws.append(["Query", job.query])
        ws.append(["Industry", job.industry or ""])
        ws.append(["Location", job.location or ""])
        ws.append(["Status", job.status.value])
        ws.append(["Started At", job.started_at or ""])
        ws.append(["Finished At", job.finished_at or ""])
        ws.append(["Discovered Count", job.discovered_count])
        ws.append(["Processed Count", job.processed_count])
        ws.append(["Valid Leads", lead_count])
        ws.append(["Duplicates", dup_count])
        ws.append(["Failures", error_count])

    def _apply_sheet_formatting(self, ws):
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        ws.freeze_panes = "A2"
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in range(2, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = thin_border

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
